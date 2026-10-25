# -*- coding: utf-8 -*-
"""Ubah docs/DOKUMENTASI_FIRMWARE.md menjadi dokumen Word untuk lampiran tesis,
lengkap dengan daftar isi dan listing kode berbaris-nomor.

Pemakaian (Python 3 dengan python-docx):
    python tools/bangun_lampiran.py

Jalur pada konstanta FW, SUMBER, dan KELUAR di bawah mengikuti mesin penulis.
SUMBER adalah dokumen Word yang gayanya (Heading 1 sampai 3, Isi, Butir, Rapat)
dipakai ulang; ubah ketiganya bila dijalankan di tempat lain. Daftar isi
diperbarui dengan membuka berkas di Word (Ctrl+A lalu F9) atau lewat COM.
Daftar berkas yang dilisting ada pada variabel DAFTAR."""

import re
import os
import docx
from docx.shared import Pt, Cm
from docx.enum.style import WD_STYLE_TYPE
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

FW = r"D:/Wahyu/TESIS/firmware/"
SUMBER = r"D:/Wahyu/TESIS/dokumen/Perbaikan REV 2 - Teks Siap Tempel - I Gede Wahyu Wiranata.docx"
KELUAR = r"D:/Wahyu/TESIS/dokumen/Lampiran Dokumentasi dan Listing Kode Firmware - I Gede Wahyu Wiranata.docx"

d = docx.Document(SUMBER)
for el in list(d.element.body):
    if el.tag != qn("w:sectPr"):
        d.element.body.remove(el)

# ---- gaya kode
kode = d.styles.add_style("KodeBaris", WD_STYLE_TYPE.PARAGRAPH)
kode.base_style = d.styles["Normal"]
kode.font.name = "Consolas"
kode.element.rPr.rFonts.set(qn("w:eastAsia"), "Consolas")
kode.element.rPr.rFonts.set(qn("w:cs"), "Consolas")
kode.font.size = Pt(7.5)
kode.paragraph_format.space_before = Pt(0)
kode.paragraph_format.space_after = Pt(0)
kode.paragraph_format.line_spacing = 1.0
kode.paragraph_format.left_indent = Cm(0.9)
kode.paragraph_format.first_line_indent = Cm(-0.9)
kode.paragraph_format.alignment = 0

kode_blok = d.styles.add_style("KodeBlok", WD_STYLE_TYPE.PARAGRAPH)
kode_blok.base_style = d.styles["Normal"]
kode_blok.font.name = "Consolas"
kode_blok.element.rPr.rFonts.set(qn("w:eastAsia"), "Consolas")
kode_blok.font.size = Pt(8)
kode_blok.paragraph_format.space_before = Pt(0)
kode_blok.paragraph_format.space_after = Pt(0)
kode_blok.paragraph_format.line_spacing = 1.0
kode_blok.paragraph_format.left_indent = Cm(0.4)


def shade(p, warna="F4F4F4"):
    ppr = p._p.get_or_add_pPr()
    sh = OxmlElement("w:shd")
    sh.set(qn("w:val"), "clear")
    sh.set(qn("w:color"), "auto")
    sh.set(qn("w:fill"), warna)
    ppr.append(sh)


def tambah_inline(p, teks, ukuran=None):
    """Pecah teks menjadi run biasa, `kode`, dan **tebal**."""
    for bag in re.split(r"(`[^`]+`|\*\*[^*]+\*\*)", teks):
        if not bag:
            continue
        if bag.startswith("`") and bag.endswith("`"):
            r = p.add_run(bag[1:-1])
            r.font.name = "Consolas"
            r._element.rPr.rFonts.set(qn("w:eastAsia"), "Consolas")
            r.font.size = Pt(ukuran - 0.5 if ukuran else 9)
        elif bag.startswith("**") and bag.endswith("**"):
            r = p.add_run(bag[2:-2])
            r.bold = True
            if ukuran:
                r.font.size = Pt(ukuran)
        else:
            r = p.add_run(bag)
            if ukuran:
                r.font.size = Pt(ukuran)


def tabel(baris, ukuran=9):
    kepala, isi = baris[0], baris[1:]
    n = len(kepala)
    t = d.add_table(rows=1, cols=n)
    b = OxmlElement("w:tblBorders")
    for sisi in ("top", "left", "bottom", "right", "insideH", "insideV"):
        e = OxmlElement(f"w:{sisi}")
        e.set(qn("w:val"), "single")
        e.set(qn("w:sz"), "4")
        e.set(qn("w:color"), "808080")
        b.append(e)
    t._tbl.tblPr.append(b)
    t.autofit = False

    def sel(c, x, tebal=False, warna=None):
        c.text = ""
        p = c.paragraphs[0]
        p.paragraph_format.space_after = Pt(2)
        p.paragraph_format.space_before = Pt(2)
        p.paragraph_format.alignment = 0
        if tebal:
            r = p.add_run(x)
            r.bold = True
            r.font.size = Pt(ukuran)
        else:
            tambah_inline(p, x, ukuran)
        if warna:
            sh = OxmlElement("w:shd")
            sh.set(qn("w:val"), "clear")
            sh.set(qn("w:color"), "auto")
            sh.set(qn("w:fill"), warna)
            c._tc.get_or_add_tcPr().append(sh)

    for k, h in enumerate(kepala):
        sel(t.rows[0].cells[k], h, True, "E8E8E8")
    for br in isi:
        cells = t.add_row().cells
        for k in range(n):
            sel(cells[k], br[k] if k < len(br) else "")
    for tr in t._tbl.tr_lst:
        tr.get_or_add_trPr().append(OxmlElement("w:cantSplit"))
    t._tbl.tr_lst[0].get_or_add_trPr().append(OxmlElement("w:tblHeader"))
    # lebar kolom sebanding dengan panjang isi, dibatasi
    bobot = []
    for k in range(n):
        m = max(len(re.sub(r"[`*]", "", r[k])) if k < len(r) else 0 for r in baris)
        bobot.append(min(max(m, 6), 55) + 6)
    total = 16.4
    for row in t.rows:
        for k, w in enumerate(bobot):
            row.cells[k].width = Cm(total * w / sum(bobot))
    d.add_paragraph(style="Rapat")


def sisip_toc():
    p = d.add_paragraph(style="Isi")
    r = p.add_run()
    for jenis, teks in (("begin", None), (None, 'TOC \\o "1-3" \\h \\z \\u'), ("separate", None), ("end", None)):
        if jenis:
            e = OxmlElement("w:fldChar")
            e.set(qn("w:fldCharType"), jenis)
        else:
            e = OxmlElement("w:instrText")
            e.set(qn("xml:space"), "preserve")
            e.text = teks
        r._r.append(e)


# ======================================================== bagian dari Markdown
md = open(FW + "docs/DOKUMENTASI_FIRMWARE.md", encoding="utf-8").read().replace("\r\n", "\n").split("\n")
i = 0
sudah_toc = False
while i < len(md):
    l = md[i]
    if l.startswith("```"):
        i += 1
        while i < len(md) and not md[i].startswith("```"):
            p = d.add_paragraph(style="KodeBlok")
            shade(p)
            p.add_run(md[i].replace("\t", "    "))
            i += 1
        d.add_paragraph(style="Rapat")
        i += 1
        continue
    if l.startswith("|"):
        baris = []
        while i < len(md) and md[i].startswith("|"):
            sel_ = [c.strip() for c in md[i].strip().strip("|").split("|")]
            if not all(re.fullmatch(r":?-{3,}:?", c) for c in sel_):
                baris.append(sel_)
            i += 1
        tabel(baris)
        continue
    if l.startswith("# "):
        d.add_paragraph(l[2:], style="Heading 1")
    elif l.startswith("## "):
        if not sudah_toc:
            sudah_toc = True
            p = d.add_paragraph(style="Isi")
            r = p.add_run("Daftar isi")
            r.bold = True
            sisip_toc()
        d.add_paragraph(l[3:], style="Heading 2")
    elif l.startswith("### "):
        d.add_paragraph(l[4:], style="Heading 3")
    elif l.startswith("- "):
        p = d.add_paragraph(style="Butir")
        tambah_inline(p, "\u2022  " + l[2:])
    elif l.strip() == "":
        pass
    else:
        p = d.add_paragraph(style="Isi")
        tambah_inline(p, l)
    i += 1

# ======================================================== listing kode
d.add_paragraph("12. Listing kode", style="Heading 2")
d.add_paragraph(
    "Listing berikut adalah isi berkas sumber apa adanya pada saat dokumen ini disusun (commit terakhir pada repositori "
    "bligede/hil-sw-ucb-vehicular-data-collection). Tiap baris diberi nomor. Berkas tidak dimodifikasi; komentar di dalam kode "
    "memuat alasan perancangan dan rujukan ke naskah, SOP, dan nomor penyimpangan.", style="Isi")

DAFTAR = [
    ("12.1 Node sensor", [
        "node/platformio.ini", "node/src/main.cpp", "node/src/config/params.h", "node/src/config/pin_map.h",
        "node/src/config/protocol.h", "node/src/modules/fsm.h", "node/src/modules/fsm.cpp",
        "node/src/modules/SW_UCB_Agent.h", "node/src/modules/SW_UCB_Agent.cpp",
        "node/src/modules/energy_meter.h", "node/src/modules/energy_meter.cpp",
        "node/src/modules/radio.h", "node/src/modules/radio.cpp",
        "node/src/modules/sensor.h", "node/src/modules/sensor.cpp",
        "node/src/modules/payload.h", "node/src/modules/payload.cpp"]),
    ("12.2 Gateway", [
        "gateway/platformio.ini", "gateway/src/main.cpp", "gateway/src/config/pin_map.h",
        "gateway/src/config/protocol.h", "gateway/src/config/secrets.example.h",
        "gateway/src/modules/unggah_sheets.h", "gateway/src/modules/unggah_sheets.cpp",
        "gateway/src/modules/tampilan_lcd.h", "gateway/src/modules/tampilan_lcd.cpp",
        "gateway/src/modules/indikator_led.h", "gateway/src/modules/indikator_led.cpp",
        "gateway/apps_script/Code.gs"]),
    ("12.3 Program uji laboratorium", [
        "node/uji/kalibrasi_ina219.cpp", "node/uji/uji_tx_arus.cpp", "node/uji/uji_tx_radio.cpp",
        "node/uji/uji_rtc_ram.cpp"]),
]


def tulis_listing(judul, baris_kode):
    d.add_paragraph(judul, style="Heading 3")
    for no, isi in enumerate(baris_kode, 1):
        p = d.add_paragraph(style="KodeBaris")
        isi = isi.replace("\t", "    ").rstrip()
        p.add_run("%4d  %s" % (no, isi))


total_baris = 0
for judul, berkas in DAFTAR:
    d.add_paragraph(judul, style="Heading 2" if False else "Heading 3")
    d.paragraphs[-1].runs[0].bold = True
    for b in berkas:
        sumber = open(FW + b, encoding="utf-8", errors="replace").read().replace("\r\n", "\n").split("\n")
        if sumber and sumber[-1] == "":
            sumber = sumber[:-1]
        total_baris += len(sumber)
        p = d.add_paragraph(style="Isi")
        r = p.add_run("%s  (%d baris)" % (b, len(sumber)))
        r.bold = True
        p.paragraph_format.keep_with_next = True
        for no, isi in enumerate(sumber, 1):
            q = d.add_paragraph(style="KodeBaris")
            q.add_run("%4d  %s" % (no, isi.replace("\t", "    ").rstrip()))

# potongan pustaka LoRa yang diubah
d.add_paragraph("12.4 Perubahan pada pustaka LoRa lokal (node/lib/LoRa)", style="Heading 3")
d.paragraphs[-1].runs[0].bold = True
lora = open(FW + "node/lib/LoRa/src/LoRa.cpp", encoding="utf-8", errors="replace").read().replace("\r\n", "\n").split("\n")
a = next(k for k, s in enumerate(lora) if "bool LoRaClass::isTransmitting()" in s)
b = next(k for k, s in enumerate(lora) if "bool LoRaClass::txDone()" in s)
akhir = next(k for k in range(b, len(lora)) if lora[k].startswith("}"))
d.add_paragraph("Potongan node/lib/LoRa/src/LoRa.cpp (baris %d sampai %d). Selebihnya adalah pustaka asli tanpa perubahan; "
                "pada LoRa.h, isTransmitting() dipindah dari private ke public dan txDone() ditambahkan." % (a + 1, akhir + 1), style="Isi")
for k in range(a, akhir + 1):
    q = d.add_paragraph(style="KodeBaris")
    q.add_run("%4d  %s" % (k + 1, lora[k].replace("\t", "    ").rstrip()))

d.add_paragraph(style="Isi")
d.add_paragraph("Alat bantu Python (folder tools/) tidak dilisting di sini; deskripsinya ada pada Bagian 10 dan sumbernya pada "
                "repositori.", style="Isi")

d.core_properties.title = "Dokumentasi dan Listing Kode Firmware Node dan Gateway"
d.core_properties.author = "I Gede Wahyu Wiranata"
d.save(KELUAR)
print("OK", KELUAR, "| baris kode terlisting:", total_baris)
